"""
Finding the tables on hand-made sheets. Each test is one layout people actually use.
Grids are lists of row tuples, empty cells as None, as the readers produce them.
"""

import datetime as dt

import openpyxl

from ingest.blocks import find_blocks
from ingest.pipeline import ingest_workbook

_ = None


def grid(*rows):
    width = max(len(r) for r in rows)
    return [tuple(r) + (None,) * (width - len(r)) for r in rows]


def headers(blocks):
    return [b.header for b in blocks]


def test_a_simple_sheet_is_one_table_with_its_title():
    g = grid(("Sales export",),
             (),
             ("region", "amount"),
             ("North", 10),
             ("South", 20))
    blocks, leftovers = find_blocks(g)
    assert headers(blocks) == [("region", "amount")]
    assert blocks[0].rows == [3, 4] and blocks[0].title == ["Sales export"]
    assert leftovers == []


def test_tables_side_by_side():
    g = grid(("city", "visits", _, "keyword", "volume"),
             ("Pune", 7, _, "sofa", 900),
             ("Delhi", 9, _, "bed", 800),
             (_, _, _, "lamp", 300))
    blocks, _left = find_blocks(g)
    assert headers(blocks) == [("city", "visits"), ("keyword", "volume")]
    assert [len(b.rows) for b in blocks] == [2, 3]


def test_stacked_tables_each_keep_their_title():
    g = grid(("Meta plan",),
             ("market", "budget"),
             ("P1", 100),
             ("P2", 200),
             (),
             ("YouTube plan",),
             ("market", "views", "budget"),
             ("P1", 5000, 300),
             ("P2", 6000, 400))
    blocks, _left = find_blocks(g)
    assert [b.title for b in blocks] == [["Meta plan"], ["YouTube plan"]]
    assert headers(blocks) == [("market", "budget"), ("market", "views", "budget")]


def test_a_block_far_to_the_right_is_its_own_table():
    far = 5000
    row = lambda *a, extra=(): tuple(a) + (None,) * (far - len(a)) + tuple(extra)
    g = [row("name", "value", extra=("label", "count")),
         row("a", 1, extra=("x", 3)),
         row("b", 2, extra=("y", 4))]
    blocks, _left = find_blocks(g)
    assert [(b.left, b.right) for b in blocks] == [(0, 2), (far, far + 2)]


def test_blank_rows_inside_a_text_heavy_table_do_not_start_a_new_table():
    g = grid(("id", "name", "city", "joined"),
             ("C1", "Asha", "Pune", dt.date(2024, 1, 5)),
             ("C2", "Ravi", "Delhi", dt.date(2024, 2, 6)),
             (),
             ("C3", "Meera", "Surat", dt.date(2024, 3, 7)),
             ("C4", "Kabir", "Goa", dt.date(2024, 4, 8)))
    blocks, _left = find_blocks(g)
    assert len(blocks) == 1
    assert blocks[0].rows == [1, 2, 4, 5] and blocks[0].skipped == [(3, "blank row")]


def test_a_key_value_preamble_is_the_title_and_footnotes_are_notes():
    g = grid(("Client", "Acme"),
             ("Duration", "Oct - Nov"),
             (),
             ("*rates are indicative",),
             ("objective", "publisher", "cost"),
             ("Traffic", "Site A", 100),
             ("Traffic", "Site B", 200))
    blocks, _left = find_blocks(g)
    assert len(blocks) == 1
    assert blocks[0].title == ["Client: Acme", "Duration: Oct - Nov"]
    assert blocks[0].notes == ["*rates are indicative"]


def test_note_rows_and_sparse_unlabelled_columns_become_notes():
    g = grid((_, "platform", "market", "budget"),
             ("raise spend in P1", "Meta", "P1", 100),
             (_, "Meta", "P2", 200),
             (_, "check the CPM with the vendor", _, _),
             (_, "Google", "P1", 300))
    b = find_blocks(g)[0][0]
    assert b.header == ("platform", "market", "budget")
    assert b.rows == [1, 2, 4]
    assert (3, "note") in b.skipped
    assert set(b.notes) == {"raise spend in P1", "check the CPM with the vendor"}


def test_an_empty_column_under_a_label_is_dropped():
    g = grid(("market", "clicks", "budget"),
             ("P1", _, 100),
             ("P2", _, 200))
    b = find_blocks(g)[0][0]
    assert b.header == ("market", "budget") and "clicks" in b.notes


def test_a_list_split_by_blank_rows_stays_one_list():
    g = grid(("Interests",), ("Sofas",), ("Beds",), ("Lamps",), (), ("Rugs",), ("Clocks",), ("Mirrors",), ("Vases",))
    blocks, _left = find_blocks(g)
    assert len(blocks) == 1 and blocks[0].header == ("Interests",) and len(blocks[0].rows) == 7


def test_a_list_under_a_label_uses_the_label_as_its_header():
    g = grid(("Furniture - 28-55",), ("Furniture",), (), ("Sofas",), ("Beds",), ("Lamps",), ("Rugs",))
    blocks, _left = find_blocks(g)
    assert blocks[0].header == ("Furniture",) and blocks[0].title == ["Furniture - 28-55"]
    assert len(blocks[0].rows) == 4


def test_a_spacer_column_does_not_split_a_table():
    g = grid(("market", "budget", _, "views"),
             ("P1", 100, _, 5000),
             ("P2", 200, _, 6000))
    blocks, _left = find_blocks(g)
    assert len(blocks) == 1 and blocks[0].header == ("market", "budget", "views")


def test_a_tiny_scrap_next_to_a_table_is_left_over_but_alone_it_is_the_data():
    scrap = grid(("market", "budget", "views"), ("P1", 100, 5000), ("P2", 200, 6000), (),
                 (_, "vendor", "rate"), (_, "OTT", 4))
    blocks, leftovers = find_blocks(scrap)
    assert len(blocks) == 1 and len(leftovers) == 1
    alone = grid(("setting", "value"), ("currency", "INR"))
    assert len(find_blocks(alone)[0]) == 1


def test_a_total_row_without_a_label_is_skipped():
    g = grid(("keyword", "volume", "clicks"),
             ("sofa", 900, 30),
             ("bed", 800, 20),
             ("lamp", 300, 10),
             (_, 2000, 60))                     # =SUM(...) with no "Total" typed next to it
    b = find_blocks(g)[0][0]
    assert b.rows == [1, 2, 3] and (4, "total / summary row") in b.skipped


def test_a_last_row_that_only_looks_like_a_total_is_kept():
    not_sums = grid(("keyword", "volume"), ("sofa", 900), ("bed", 800), (_, 1000))
    assert find_blocks(not_sums)[0][0].rows == [1, 2, 3]
    labelled = grid(("keyword", "volume"), ("sofa", 900), ("bed", 800), ("lamp", 1700))
    assert find_blocks(labelled)[0][0].rows == [1, 2, 3]


def test_a_grid_with_no_header_is_all_leftover():
    g = grid(("P1", "Pune", 411001), ("P2", "Surat", 395007), ("P3", "Goa", 403001))
    blocks, leftovers = find_blocks(g)
    assert blocks == [] and len(leftovers) == 1


def test_sheet_tables_get_names_ranges_and_titles(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plan"
    ws["B2"] = "Meta RnF"
    for r, row in enumerate([("market", "budget"), ("P1", 100), ("P2", 200)], start=3):
        for c, v in enumerate(row, start=2):
            ws.cell(r, c, v)
    ws["B7"] = "*budgets are estimates"
    ws["B8"] = "YouTube"
    for r, row in enumerate([("market", "views"), ("P1", 5000), ("P2", 6000)], start=9):
        for c, v in enumerate(row, start=2):
            ws.cell(r, c, v)
    wb.save(tmp_path / "wb.xlsx")
    reports = ingest_workbook(tmp_path / "wb.xlsx")
    assert [(r.table, r.range, r.title) for r in reports] == [
        ("plan_meta_rnf", "B2:C5", "Meta RnF"),
        ("plan_youtube", "B7:C11", "YouTube"),
    ]
    assert reports[1].notes == ["*budgets are estimates"]
    assert [r.rows_loaded for r in reports] == [2, 2]
