"""
One bad sheet must not sink the whole workbook.

A sheet that is too wide for Postgres is skipped with a reason before any SQL runs,
and a sheet that fails inside the database is rolled back to its own savepoint
while the other sheets still load. The database test needs scripts/setup_db.py.
"""

import os
from pathlib import Path

import openpyxl
import psycopg
import pytest
from dotenv import load_dotenv

from ingest import pipeline
from ingest.pipeline import Sink, ingest_workbook

load_dotenv(Path(__file__).resolve().parents[2] / ".env")
APP_URL = os.environ.get("DATABASE_URL")


def _workbook(path, wide_columns=2000):
    """Two good sheets with a 'Wide' sheet between them, whose table has `wide_columns` columns."""
    wb = openpyxl.Workbook()
    good = wb.active
    good.title = "Good"
    good.append(["region", "amount"])
    good.append(["North", 100])
    good.append(["South", 250])
    wide = wb.create_sheet("Wide")
    wide.append([f"c{i}" for i in range(wide_columns)])
    wide.append(list(range(wide_columns)))
    wide.append(list(range(wide_columns)))
    last = wb.create_sheet("Also good")
    last.append(["city", "visits"])
    last.append(["Pune", 7])
    wb.save(path)
    return path


def _by_sheet(reports):
    return {r.sheet: r for r in reports}


def test_too_wide_sheet_is_skipped_with_a_reason(tmp_path):
    reports = _by_sheet(ingest_workbook(_workbook(tmp_path / "wb.xlsx")))
    assert reports["Wide"].status == "skipped"
    assert "too wide: 2,000 columns" in reports["Wide"].reason
    assert reports["Good"].status == reports["Also good"].status == "loaded"


def test_a_sheet_that_fails_in_the_sink_is_skipped_and_the_rest_load(tmp_path):
    class FailsOnWide(Sink):
        aborted = []

        def begin_table(self, table, columns):
            if table == "wide":
                raise RuntimeError("disk on fire\nsecond line")

        def abort_table(self, exc):
            self.aborted.append(str(exc))

    sink = FailsOnWide()
    reports = _by_sheet(ingest_workbook(_workbook(tmp_path / "wb.xlsx", wide_columns=3), sink))
    assert reports["Wide"].status == "skipped"
    assert reports["Wide"].reason == "couldn't be loaded: disk on fire"
    assert reports["Wide"].table is None and reports["Wide"].rows_loaded == 0
    assert sink.aborted == ["disk on fire\nsecond line"]
    assert reports["Good"].rows_loaded == 2 and reports["Also good"].rows_loaded == 1


@pytest.mark.skipif(not APP_URL, reason="run scripts/setup_db.py first")
def test_a_sheet_postgres_rejects_rolls_back_alone(tmp_path, monkeypatch):
    from ingest.postgres import load_workbook

    # Let the wide sheet through the guard, so Postgres itself rejects the CREATE TABLE.
    monkeypatch.setattr(pipeline, "MAX_COLUMNS", 10_000)
    schema = "ds_isolation_test"
    with psycopg.connect(APP_URL, autocommit=True) as conn:
        conn.execute(f"DROP SCHEMA IF EXISTS {schema} CASCADE")
    try:
        _, reports, _ = load_workbook(_workbook(tmp_path / "wb.xlsx"), dsn=APP_URL, schema=schema)
        reports = _by_sheet(reports)
        assert reports["Wide"].status == "skipped"
        assert "1600 columns" in reports["Wide"].reason
        with psycopg.connect(APP_URL) as conn:
            tables = {r[0] for r in conn.execute(
                "SELECT table_name FROM information_schema.tables WHERE table_schema = %s", (schema,))}
            assert tables == {"good", "also_good"}
            assert conn.execute(f"SELECT sum(amount) FROM {schema}.good").fetchone()[0] == 350
    finally:
        with psycopg.connect(APP_URL, autocommit=True) as conn:
            conn.execute(f"DROP SCHEMA IF EXISTS {schema} CASCADE")
