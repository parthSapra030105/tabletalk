import openpyxl

from ingest.pipeline import ingest_workbook
from ingest.readers import WorkbookReader, _fill_merged_down
from ingest.values import is_null


def _plan_sheet(path):
    """A small hand-made plan: starts at B2, a title merged across, a category merged down."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plan"
    ws["B2"] = "Summary"
    ws.merge_cells("B2:E2")
    ws.append([])
    for row in (["Category", "Platform", "Budget", "Note"],
                ["Furniture", "Meta", 100, "   "],
                [None, "YouTube", 200, "\xa0"],
                [None, "Google", 300, None],
                ["Decor", "Meta", 50, "_"],
                [None, "YouTube", 70, None]):
        ws.append([None, *row])
    ws.merge_cells("B5:B7")
    ws.merge_cells("B8:B9")
    wb.save(path)
    return path


def test_rows_are_normalised_and_merges_filled_down(tmp_path):
    reader = WorkbookReader(_plan_sheet(tmp_path / "wb.xlsx"))
    try:
        rows = list(reader.rows(reader.sheets()[0]))
    finally:
        reader.close()
    # calamine yields rows from Excel row 1 but columns from the first used one (B), so column 0 is B.
    assert rows[0] == (None, None, None, None)
    assert rows[1] == ("Summary", None, None, None)          # a title merged across is not repeated
    assert [r[0] for r in rows[4:]] == ["Furniture"] * 3 + ["Decor"] * 2
    assert rows[4][3] is None and rows[5][3] is None          # whitespace and a non-breaking space are empty


def test_title_merge_does_not_become_the_header(tmp_path):
    report = ingest_workbook(_plan_sheet(tmp_path / "wb.xlsx"))[0]
    assert report.header_row == 4
    assert [c.name for c in report.columns][:3] == ["category", "platform", "budget"]
    assert report.columns[2].type == "numeric"
    assert report.rows_loaded == 5


def test_fill_down_with_an_offset_origin():
    rows = [("a", 1), (None, 2), (None, 3), ("b", 4)]
    # Absolute rows 5-7, column 2 merged; the rows start at absolute (5, 2).
    merges = [((5, 2), (7, 2))]
    assert [r[0] for r in _fill_merged_down(iter(rows), merges, (5, 2))] == ["a", "a", "a", "b"]


def test_fill_down_keeps_a_value_already_in_the_cell():
    rows = [("a",), ("kept",)]
    assert list(_fill_merged_down(iter(rows), [((0, 0), (1, 0))], (0, 0))) == [("a",), ("kept",)]


def test_underscore_is_a_null_token():
    assert is_null("_") and is_null(" _ ")
